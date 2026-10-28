"""Presupuestos: armador de cotizaciones y generación del PDF."""

import logging
import uuid
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, File, Form, HTTPException, Response, UploadFile, status
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from app.api.deps import get_current_user, get_user_gmail, resolver_duenio
from app.api.listing import scalars_capped
from app.config import get_settings
from app.core.exceptions import NotFoundError
from app.db.models.oportunidades import AmbitoOportunidad, Oportunidad
from app.db.models.presupuesto_items import PresupuestoItem
from app.db.models.presupuestos import SUCURSALES_GBP, EstadoCargaGbp, Presupuesto
from app.db.models.solicitudes_compras import SolicitudCompras
from app.db.models.usuarios import Usuario
from app.db.session import get_db
from app.integrations.gmail.errors import clasificar_error_gmail
from app.schemas.presupuesto import (
    EnviarPresupuestoRequest,
    ItemBase,
    PresupuestoCreate,
    PresupuestoListItem,
    PresupuestoRead,
    PresupuestoUpdate,
)
from app.services import presupuestos_gbp as gbp_svc
from app.services.importar_excel import parsear_y_validar
from app.services.presupuestos import (
    actualizar_presupuesto,
    crear_desde_solicitud,
    crear_lote,
    crear_presupuesto,
    enviar_al_cliente,
)
from app.services.storage import get_storage

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/presupuestos", tags=["presupuestos"])

_RELATIONS = (
    selectinload(Presupuesto.items),
    selectinload(Presupuesto.oportunidad).selectinload(Oportunidad.cliente),
    selectinload(Presupuesto.oportunidad).selectinload(Oportunidad.contacto),
    selectinload(Presupuesto.creado_por),
    selectinload(Presupuesto.editado_por),
)

# Listado: sin ítems (el mayor peso por fila) ni editado_por; el detalle usa su
# propio fetch con todo.
_LIST_RELATIONS = (
    selectinload(Presupuesto.oportunidad).selectinload(Oportunidad.cliente),
    selectinload(Presupuesto.creado_por),
)


def _get_loaded(db: Session, presupuesto_id: int) -> Presupuesto:
    presupuesto = db.get(Presupuesto, presupuesto_id, options=list(_RELATIONS))
    if presupuesto is None:
        raise NotFoundError("Presupuesto no encontrado")
    return presupuesto


def _assert_owner(presupuesto: Presupuesto, user: Usuario) -> None:
    """Acceso compartido: cualquier usuario puede ver/editar cualquier
    presupuesto. Se mantiene (no-op) para no tocar los call sites."""
    return


@router.get("", response_model=list[PresupuestoListItem])
def list_presupuestos(
    oportunidad_id: int | None = None,
    usuario_id: int | None = None,
    ambito: AmbitoOportunidad | None = None,
    con_omitidos: bool = False,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
) -> list[Presupuesto]:
    """Presupuestos personales: cada vendedor ve solo los de sus oportunidades.
    Los admin ven los de todo el equipo, o filtran por `usuario_id` (perfil).
    `ambito` limita a la sección (por el ámbito de la oportunidad del presupuesto)."""
    query = select(Presupuesto).options(*_LIST_RELATIONS).order_by(Presupuesto.id.desc())
    if oportunidad_id is not None:
        query = query.where(Presupuesto.oportunidad_id == oportunidad_id)
    duenio_id = resolver_duenio(current_user, usuario_id)
    if duenio_id is not None or ambito is not None:
        # OUTER join: los presupuestos STANDALONE (sin oportunidad, p.ej. los que
        # vienen de un Excel) no tienen ámbito ni vendedor, pero igual tienen que
        # verse. Por eso se incluyen aunque haya filtro de sección/dueño.
        sin_op = Presupuesto.oportunidad_id.is_(None)
        query = query.outerjoin(Oportunidad, Presupuesto.oportunidad_id == Oportunidad.id)
        if ambito is not None:
            query = query.where((Oportunidad.ambito == ambito.value) | sin_op)
        if duenio_id is not None:
            query = query.where(
                (Oportunidad.vendedor_id == duenio_id)
                | (sin_op & (Presupuesto.creado_por_id == duenio_id))
            )
    presupuestos = scalars_capped(db, query, "presupuestos")

    # Totales reales por moneda (los ítems pueden ser ARS y/o USD; sumarlos en una
    # sola cifra con la moneda del documento engaña). Una agregación aparte evita
    # traer todos los ítems al listado.
    ids = [p.id for p in presupuestos]
    totales: dict[int, dict[str, object]] = {}
    if ids:
        filas = db.execute(
            select(
                PresupuestoItem.presupuesto_id,
                PresupuestoItem.moneda,
                func.sum(PresupuestoItem.cantidad * PresupuestoItem.precio_unitario),
            )
            .where(PresupuestoItem.presupuesto_id.in_(ids))
            .group_by(PresupuestoItem.presupuesto_id, PresupuestoItem.moneda)
        )
        for pid, moneda, total in filas:
            totales.setdefault(pid, {})[moneda or "USD"] = total

    salida = []
    for p in presupuestos:
        cant_omitidos = len(p.items_omitidos or [])
        if con_omitidos and cant_omitidos == 0:
            continue
        item = PresupuestoListItem.model_validate(p)
        item.totales_por_moneda = totales.get(p.id, {})
        item.cant_omitidos = cant_omitidos
        salida.append(item)
    return salida


@router.post("", response_model=PresupuestoRead, status_code=201)
def create_presupuesto(
    body: PresupuestoCreate,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
) -> Presupuesto:
    """Crea un presupuesto (borrador). Con oportunidad la avanza a 'presupuestada';
    standalone (solo N° de cliente) si no se pasa oportunidad."""
    if body.oportunidad_id is not None and db.get(Oportunidad, body.oportunidad_id) is None:
        raise HTTPException(status_code=404, detail="La oportunidad no existe")
    presupuesto = crear_presupuesto(db, body)
    presupuesto.creado_por_id = current_user.id
    db.commit()
    return _get_loaded(db, presupuesto.id)


@router.get("/sucursales")
def listar_sucursales(_: Usuario = Depends(get_current_user)) -> list[dict[str, str]]:
    """Opciones de sucursal de GBP (fuente única: SUCURSALES_GBP)."""
    return [{"valor": v, "nombre": n} for v, n in SUCURSALES_GBP.items()]


@router.post("/lote", response_model=list[PresupuestoRead], status_code=201)
def create_lote(
    body: PresupuestoCreate,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
) -> list[Presupuesto]:
    """Crea el presupuesto; si tiene más de GBP_MAX_ITEMS ítems lo parte en varias
    'partes' (cada una = un pedido de GBP). Devuelve la lista (1 o N)."""
    if body.oportunidad_id is not None and db.get(Oportunidad, body.oportunidad_id) is None:
        raise HTTPException(status_code=404, detail="La oportunidad no existe")
    creados = crear_lote(db, body, current_user.id)
    return [_get_loaded(db, p.id) for p in creados]


@router.post("/grupo/{grupo}/cargar-gbp", response_model=list[PresupuestoRead])
def cargar_grupo_en_gbp(
    grupo: str,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
) -> list[Presupuesto]:
    """Encola TODAS las partes de un lote para que el bot las cargue en GBP."""
    partes = list(
        db.scalars(
            select(Presupuesto)
            .where(Presupuesto.gbp_grupo == grupo)
            .order_by(Presupuesto.gbp_parte)
            .options(*_RELATIONS)
        )
    )
    if not partes:
        raise HTTPException(status_code=404, detail="Grupo no encontrado")
    try:
        for p in partes:
            try:
                gbp_svc.encolar(db, p)
            except (
                gbp_svc.YaCargadoError,
                gbp_svc.RequiereConfirmacionError,
                gbp_svc.CanceladoError,
            ):
                continue  # cargado/'revisar'/cancelado: no se toca (se reencola por parte)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
    return [_get_loaded(db, p.id) for p in partes]


def _omitidos_del_grupo(db: Session, grupo: str) -> list[dict]:
    """Vista consolidada de los ítems omitidos de todas las partes de un lote."""
    partes = db.scalars(
        select(Presupuesto)
        .where(Presupuesto.gbp_grupo == grupo)
        .order_by(Presupuesto.gbp_parte)
    ).all()
    filas: list[dict] = []
    for p in partes:
        for it in p.items_omitidos or []:
            filas.append(
                {
                    "parte": p.gbp_parte,
                    "numero_gbp": p.id_gbp_pedido,
                    "linea": it.get("linea"),
                    "sku": it.get("sku"),
                    "descripcion": it.get("descripcion"),
                    "motivo": it.get("motivo"),
                }
            )
    return filas


@router.get("/grupo/{grupo}/omitidos")
def omitidos_grupo(
    grupo: str,
    db: Session = Depends(get_db),
    _: Usuario = Depends(get_current_user),
) -> list[dict]:
    """Ítems no cargados de todas las partes del lote (SKU + parte + N° de GBP)."""
    return _omitidos_del_grupo(db, grupo)


@router.get("/grupo/{grupo}/omitidos.xlsx")
def exportar_omitidos_grupo(
    grupo: str,
    db: Session = Depends(get_db),
    _: Usuario = Depends(get_current_user),
) -> Response:
    """Exporta a Excel los ítems no cargados del lote."""
    from io import BytesIO

    from openpyxl import Workbook

    filas = _omitidos_del_grupo(db, grupo)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ítems no cargados"
    ws.append(["Parte", "N° GBP", "Línea", "SKU", "Descripción", "Motivo"])
    for f in filas:
        ws.append(
            [f["parte"], f["numero_gbp"], f["linea"], f["sku"], f["descripcion"], f["motivo"]]
        )
    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="omitidos-{grupo}.xlsx"'},
    )


@router.get("/grupo/{grupo}/resultado.xlsx")
def exportar_resultado_grupo(
    grupo: str,
    db: Session = Depends(get_db),
    _: Usuario = Depends(get_current_user),
) -> Response:
    """Devuelve el Excel ORIGINAL del lote con el resultado de GBP: 3 columnas
    nuevas (Estado GBP / Presupuesto GBP / Motivo), filas de ítems no cargados en
    naranja y filas de partes en error en rojo claro. Se puede bajar en cualquier
    momento (refleja lo que haya hasta ahí)."""
    from io import BytesIO

    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    partes = list(
        db.scalars(
            select(Presupuesto)
            .where(Presupuesto.gbp_grupo == grupo)
            .order_by(Presupuesto.gbp_parte)
            .options(selectinload(Presupuesto.items))
        )
    )
    if not partes:
        raise HTTPException(status_code=404, detail="Lote no encontrado")
    xlsx_key = next((p.lote_xlsx_key for p in partes if p.lote_xlsx_key), None)
    if not xlsx_key:
        raise HTTPException(
            status_code=404, detail="Este lote no tiene el Excel original guardado."
        )
    nombre = next((p.lote_nombre for p in partes if p.lote_nombre), None) or f"lote-{grupo}"

    # fila_excel -> (estado, numero_gbp, motivo, color)
    naranja = PatternFill(fill_type="solid", fgColor="FFC000")
    rojo = PatternFill(fill_type="solid", fgColor="F4CCCC")
    _ESTADO = {
        "cargado": "Cargado",
        "error": "Error",
        "pendiente": "Pendiente",
        "en_proceso": "Pendiente",
        "cancelado": "Cancelado",
        None: "Pendiente",
    }
    filas_res: dict[int, tuple] = {}
    for p in partes:
        omit = {o.get("linea"): o.get("motivo") for o in (p.items_omitidos or [])}
        for pos, item in enumerate(p.items, start=1):  # pos = "linea" (1-based)
            if item.fila_excel is None:
                continue
            if p.estado_carga == "error":
                filas_res[item.fila_excel] = ("Error", p.id_gbp_pedido, p.gbp_error, rojo)
            elif pos in omit:
                filas_res[item.fila_excel] = (
                    "NO cargado", p.id_gbp_pedido, omit[pos], naranja,
                )
            else:
                filas_res[item.fila_excel] = (
                    _ESTADO.get(p.estado_carga, "Pendiente"), p.id_gbp_pedido, None, None,
                )

    data = get_storage().get(xlsx_key)
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    base = ws.max_column
    for off, titulo in enumerate(("Estado GBP", "Presupuesto GBP", "Motivo"), start=1):
        ws.cell(row=1, column=base + off, value=titulo)
    ultima = base + 3
    for fila, (estado, numero, motivo, fill) in filas_res.items():
        ws.cell(row=fila, column=base + 1, value=estado)
        ws.cell(row=fila, column=base + 2, value=numero)
        ws.cell(row=fila, column=base + 3, value=motivo)
        if fill is not None:
            for col in range(1, ultima + 1):
                ws.cell(row=fila, column=col).fill = fill

    buf = BytesIO()
    wb.save(buf)
    base_nombre = nombre[:-5] if nombre.lower().endswith(".xlsx") else nombre
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{base_nombre}_resultado_GBP.xlsx"'
        },
    )


@router.post("/importar-lote", response_model=list[PresupuestoRead], status_code=201)
async def importar_lote(
    archivo: UploadFile = File(...),
    cliente_numero: str = Form(...),
    entrega_por: str = Form(...),
    documento: str = Form("Presupuesto USD"),
    sucursal: str = Form("1"),
    cliente_id: int | None = Form(None),
    oportunidad_id: int | None = Form(None),
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
) -> list[Presupuesto]:
    """Sube el Excel, valida TODO el archivo y, si no hay errores, crea el lote
    (partes de GBP_MAX_ITEMS). Si hay errores, devuelve 422 con la lista
    (fila + mensaje) y NO crea/encola nada."""
    if sucursal not in SUCURSALES_GBP:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Sucursal inválida. Opciones: {', '.join(SUCURSALES_GBP)}",
        )
    contenido = await archivo.read()
    try:
        items, errores = parsear_y_validar(
            contenido, umbral_usd=get_settings().GBP_PRECIO_USD_MAX
        )
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
    if errores:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={
                "mensaje": f"{len(errores)} fila(s) con errores. No se encoló nada.",
                "errores": errores,
            },
        )
    if oportunidad_id is not None and db.get(Oportunidad, oportunidad_id) is None:
        raise HTTPException(status_code=404, detail="La oportunidad no existe")
    data = PresupuestoCreate(
        oportunidad_id=oportunidad_id,
        cliente_id=cliente_id,
        cliente_numero=cliente_numero,
        entrega_por=entrega_por,
        documento=documento,
        sucursal=sucursal,
        items=[ItemBase.model_validate(it) for it in items],
    )
    # Guardamos el .xlsx original del lote para poder regenerar el "Excel con
    # resultado" (resaltando los ítems que GBP no cargó).
    grupo = uuid.uuid4().hex[:12]
    xlsx_key = f"presupuestos_gbp/lote/{grupo}.xlsx"
    get_storage().put(
        xlsx_key,
        contenido,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    creados = crear_lote(
        db,
        data,
        current_user.id,
        grupo=grupo,
        xlsx_key=xlsx_key,
        nombre_archivo=archivo.filename,
    )
    return [_get_loaded(db, p.id) for p in creados]


@router.post("/desde-solicitud/{solicitud_id}", response_model=PresupuestoRead, status_code=201)
def create_desde_solicitud(
    solicitud_id: int,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
) -> Presupuesto:
    """Crea un presupuesto pre-llenado con la respuesta de Compras ya parseada."""
    solicitud = db.get(
        SolicitudCompras,
        solicitud_id,
        options=[selectinload(SolicitudCompras.respuestas)],
    )
    if solicitud is None:
        raise HTTPException(status_code=404, detail="La solicitud no existe")
    try:
        presupuesto = crear_desde_solicitud(db, solicitud)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
    except ValidationError as exc:
        # La respuesta de Compras guardada no coincide con el schema actual
        # (datos viejos/incompletos): mensaje claro en vez de un 500 opaco.
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                "La respuesta de Compras guardada está incompleta o desactualizada. "
                "Abrí el detalle de la solicitud y volvé a parsear la respuesta."
            ),
        ) from exc
    presupuesto.creado_por_id = current_user.id
    db.commit()
    return _get_loaded(db, presupuesto.id)


@router.get("/{presupuesto_id}", response_model=PresupuestoRead)
def get_presupuesto(
    presupuesto_id: int,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
) -> Presupuesto:
    presupuesto = _get_loaded(db, presupuesto_id)
    _assert_owner(presupuesto, current_user)
    return presupuesto


@router.patch("/{presupuesto_id}", response_model=PresupuestoRead)
def update_presupuesto(
    presupuesto_id: int,
    body: PresupuestoUpdate,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
) -> Presupuesto:
    presupuesto = _get_loaded(db, presupuesto_id)
    _assert_owner(presupuesto, current_user)
    actualizar_presupuesto(db, presupuesto, body)
    presupuesto.editado_por_id = current_user.id
    presupuesto.editado_en = datetime.now(timezone.utc)
    db.commit()
    return _get_loaded(db, presupuesto_id)


@router.delete("/{presupuesto_id}", status_code=204)
def delete_presupuesto(
    presupuesto_id: int,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
) -> Response:
    presupuesto = _get_loaded(db, presupuesto_id)
    _assert_owner(presupuesto, current_user)
    # Si está en la cola del bot (pendiente/en_proceso), NO se borra de verdad: se
    # marca "cancelado" (terminal). Así el bot lo ignora, el reciclado no lo vuelve
    # a encolar, y si el bot ya lo estaba cargando podemos capturar su resultado
    # (ver registrar_resultado). Para el resto (nunca encolado o ya terminado),
    # borrado normal.
    en_cola = presupuesto.estado_carga in (
        EstadoCargaGbp.pendiente.value,
        EstadoCargaGbp.en_proceso.value,
    )
    if en_cola:
        presupuesto.estado_carga = EstadoCargaGbp.cancelado.value
        db.commit()
        return Response(status_code=204)
    for item in list(presupuesto.items):
        db.delete(item)
    db.delete(presupuesto)
    db.commit()
    return Response(status_code=204)


@router.post("/{presupuesto_id}/enviar", response_model=PresupuestoRead)
def enviar_presupuesto(
    presupuesto_id: int,
    body: EnviarPresupuestoRequest,
    db: Session = Depends(get_db),
    gmail=Depends(get_user_gmail),  # noqa: ANN001 - GmailClient del usuario logueado
    current_user: Usuario = Depends(get_current_user),
) -> Presupuesto:
    """Envía el presupuesto (PDF) al cliente por Gmail. Si no se indica `to`,
    usa el email del contacto de la oportunidad."""
    presupuesto = _get_loaded(db, presupuesto_id)
    _assert_owner(presupuesto, current_user)
    contacto = presupuesto.oportunidad.contacto if presupuesto.oportunidad else None
    destino = (body.to and str(body.to)) or (contacto.email if contacto else None)
    if not destino:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No hay email del cliente. Indicá uno o cargá el email del contacto.",
        )
    try:
        enviar_al_cliente(
            db, gmail, presupuesto, to=destino, mensaje=body.mensaje, remitente=current_user.email
        )
    except RuntimeError as exc:  # weasyprint / PDF no disponible
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail=str(exc)
        ) from exc
    except Exception as exc:  # noqa: BLE001 - frontera con Gmail
        logger.warning(
            "Falló enviar presupuesto (id=%s, usuario=%s): %s",
            presupuesto_id,
            current_user.id,
            exc,
        )
        code, msg = clasificar_error_gmail(exc)
        raise HTTPException(status_code=code, detail=msg) from exc
    return _get_loaded(db, presupuesto_id)


@router.post("/{presupuesto_id}/cargar-gbp", response_model=PresupuestoRead)
def cargar_en_gbp(
    presupuesto_id: int,
    confirmar: bool = False,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
) -> Presupuesto:
    """Encola la parte para el bot ('Cargar en GBP' / reencolar por parte).
    Rechaza si ya está 'cargado' (409) y, si quedó en 'revisar' (pudo haber
    quedado guardado en GBP), exige `confirmar=true` para no duplicar."""
    presupuesto = _get_loaded(db, presupuesto_id)
    _assert_owner(presupuesto, current_user)
    try:
        gbp_svc.encolar(db, presupuesto, confirmar=confirmar)
    except (gbp_svc.YaCargadoError, gbp_svc.CanceladoError) as exc:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc
    except gbp_svc.RequiereConfirmacionError as exc:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
    return _get_loaded(db, presupuesto_id)


@router.get("/{presupuesto_id}/pdf")
def descargar_pdf(
    presupuesto_id: int,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
) -> Response:
    """Devuelve el PDF OFICIAL del presupuesto, que genera GBP. Si todavía no se
    cargó en GBP, no hay PDF (404): el documento sale del ERP, no del CRM."""
    presupuesto = _get_loaded(db, presupuesto_id)
    _assert_owner(presupuesto, current_user)
    if not presupuesto.pdf_gbp_key:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Todavía no hay PDF: el presupuesto aún no se cargó en GBP.",
        )
    data = get_storage().get(presupuesto.pdf_gbp_key)
    return Response(
        content=data,
        media_type="application/pdf",
        headers={"Content-Disposition": f'inline; filename="{presupuesto.codigo}.pdf"'},
    )
