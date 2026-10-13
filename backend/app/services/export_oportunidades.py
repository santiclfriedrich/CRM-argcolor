"""Exportación de oportunidades a Excel (.xlsx) con openpyxl.

Replica las columnas de las tablas del CRM. En Gubernamental suma las columnas
de licitación (proceso, expediente, apertura, etc.)."""

import io
from datetime import date, datetime, time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.db.models.oportunidades import Oportunidad

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _fecha(valor: date | datetime | None) -> str:
    if valor is None:
        return ""
    return valor.strftime("%d/%m/%Y")


def _hora(valor: time | None) -> str:
    return valor.strftime("%H:%M") if valor is not None else ""


def _estado(valor: object) -> str:
    # El estado puede venir como enum o como str; lo humanizamos.
    crudo = getattr(valor, "value", valor)
    return str(crudo or "").replace("_", " ").capitalize()


# (encabezado, función que saca el valor de la oportunidad)
_COLS_BASE: list[tuple[str, object]] = [
    ("N°", lambda o: o.id),
    ("Cliente", lambda o: o.cliente.razon_social if o.cliente else ""),
    ("N° cliente", lambda o: (o.cliente.numero_cliente if o.cliente else "") or ""),
    ("Vendedor", lambda o: o.vendedor.nombre if o.vendedor else ""),
    ("Estado", lambda o: _estado(o.estado)),
    ("Asunto", lambda o: o.asunto or ""),
    ("Rubro / Producto", lambda o: o.producto or ""),
    ("N° pedido", lambda o: o.numero_pedido or ""),
    ("Valor estimado", lambda o: float(o.valor_estimado) if o.valor_estimado is not None else None),
    ("Cargada en GBP", lambda o: "Sí" if o.cargada_en_gbp else "No"),
    ("Observación", lambda o: o.observacion or ""),
    ("Fecha creación", lambda o: _fecha(o.fecha_creacion)),
    ("Pedido del cliente", lambda o: _fecha(o.fecha_pedido_cliente)),
    ("Enviado a compras", lambda o: _fecha(o.fecha_enviado_compras)),
    ("Respuesta de compras", lambda o: _fecha(o.fecha_respuesta_compras)),
    ("Enviado al cliente", lambda o: _fecha(o.fecha_enviado_cliente)),
    ("Fecha límite", lambda o: _fecha(o.fecha_limite)),
    ("Fecha entrega", lambda o: _fecha(o.fecha_entrega)),
]

_COLS_GUB: list[tuple[str, object]] = [
    ("Proceso", lambda o: o.proceso or ""),
    ("Expediente", lambda o: o.expediente or ""),
    ("Portal", lambda o: o.portal or ""),
    ("Apertura", lambda o: _fecha(o.apertura)),
    ("Hora pliego", lambda o: _hora(o.hr_pliego)),
    ("Hora apertura", lambda o: _hora(o.hr_apertura)),
    ("Moneda", lambda o: o.moneda or ""),
    ("Pliego", lambda o: o.pliego or ""),
    ("Empresa", lambda o: o.empresa or ""),
    ("Días", lambda o: o.dias),
    ("Link presupuesto", lambda o: o.presupuesto_url or ""),
]


def oportunidades_a_xlsx(
    oportunidades: list[Oportunidad], *, incluir_gub: bool
) -> bytes:
    """Arma un .xlsx con las oportunidades y lo devuelve como bytes."""
    cols = _COLS_BASE + (_COLS_GUB if incluir_gub else [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Oportunidades"

    header_fill = PatternFill("solid", fgColor="1E272C")  # antracita
    header_font = Font(bold=True, color="FFFFFF")
    for idx, (titulo, _) in enumerate(cols, start=1):
        celda = ws.cell(row=1, column=idx, value=titulo)
        celda.fill = header_fill
        celda.font = header_font
        celda.alignment = Alignment(vertical="center")

    for fila, op in enumerate(oportunidades, start=2):
        for col, (_, getter) in enumerate(cols, start=1):
            ws.cell(row=fila, column=col, value=getter(op))  # type: ignore[operator]

    # Ancho de columna aproximado según el contenido (acotado).
    for idx, (titulo, getter) in enumerate(cols, start=1):
        ancho = len(titulo)
        for op in oportunidades:
            valor = getter(op)  # type: ignore[operator]
            if valor is not None:
                ancho = max(ancho, len(str(valor)))
        ws.column_dimensions[get_column_letter(idx)].width = min(max(ancho + 2, 10), 45)

    ws.freeze_panes = "A2"  # fija la fila de encabezados al scrollear

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
