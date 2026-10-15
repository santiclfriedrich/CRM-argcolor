"""Exportación de cuentas/clientes a Excel (.xlsx) con openpyxl."""

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.db.models.clientes import Cliente

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _fecha(valor: date | datetime | None) -> str:
    return valor.strftime("%d/%m/%Y") if valor is not None else ""


def _contacto(c: Cliente):
    """Contacto representativo: el principal con mail; si no, el primero con mail;
    si no, el principal; si no, el primero. Puede ser None."""
    contactos = list(getattr(c, "contactos", None) or [])
    if not contactos:
        return None
    return (
        next((k for k in contactos if k.es_principal and k.email), None)
        or next((k for k in contactos if k.email), None)
        or next((k for k in contactos if k.es_principal), None)
        or contactos[0]
    )


def _emails(c: Cliente) -> str:
    """Todos los mails de los contactos, sin repetir."""
    vistos: list[str] = []
    for k in getattr(c, "contactos", None) or []:
        if k.email and k.email not in vistos:
            vistos.append(k.email)
    return ", ".join(vistos)


# (encabezado, función que saca el valor del cliente)
_COLS: list[tuple[str, object]] = [
    ("N° cliente", lambda c: c.numero_cliente or ""),
    ("Razón social", lambda c: c.razon_social),
    ("CUIT", lambda c: c.cuit or ""),
    ("Tipo", lambda c: c.tipo or ""),
    ("Sector", lambda c: c.sector or ""),
    ("Teléfono", lambda c: c.telefono or ""),
    ("Contacto", lambda c: (_contacto(c).nombre if _contacto(c) else "")),
    ("Email", lambda c: (_contacto(c).email or "" if _contacto(c) else "")),
    ("Emails (todos)", _emails),
    ("Sitio web", lambda c: c.sitio_web or ""),
    ("Empleados", lambda c: c.empleados),
    ("Dirección facturación", lambda c: c.direccion_facturacion or ""),
    ("Dirección envío", lambda c: c.direccion_envio or ""),
    ("Vendedor asignado", lambda c: c.vendedor_asignado.nombre if c.vendedor_asignado else ""),
    ("Activo", lambda c: "Sí" if c.activo else "No"),
    ("Notas", lambda c: c.notas or ""),
    ("Alta", lambda c: _fecha(getattr(c, "created_at", None))),
]


def clientes_a_xlsx(clientes: list[Cliente]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cuentas"

    header_fill = PatternFill("solid", fgColor="1E272C")  # antracita
    header_font = Font(bold=True, color="FFFFFF")
    for idx, (titulo, _) in enumerate(_COLS, start=1):
        celda = ws.cell(row=1, column=idx, value=titulo)
        celda.fill = header_fill
        celda.font = header_font
        celda.alignment = Alignment(vertical="center")

    for fila, cli in enumerate(clientes, start=2):
        for col, (_, getter) in enumerate(_COLS, start=1):
            ws.cell(row=fila, column=col, value=getter(cli))  # type: ignore[operator]

    for idx, (titulo, getter) in enumerate(_COLS, start=1):
        ancho = len(titulo)
        for cli in clientes:
            valor = getter(cli)  # type: ignore[operator]
            if valor is not None:
                ancho = max(ancho, len(str(valor)))
        ws.column_dimensions[get_column_letter(idx)].width = min(max(ancho + 2, 10), 50)

    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
