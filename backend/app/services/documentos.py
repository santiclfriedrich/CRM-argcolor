"""Lectura de adjuntos-documento de un mail para alimentar a la IA.

Dos tipos de adjunto útiles al cotizar:
- PDF: se pasa tal cual a la IA multimodal (Gemini lo lee nativo, tablas incluidas).
- Planilla (Excel .xlsx / CSV): se convierte a texto y se anexa al cuerpo del mail,
  porque en los RFQ los ítems a cotizar suelen venir en la planilla, no en el texto.

Todo es defensivo: un adjunto ilegible nunca corta la ingesta (devuelve None).
"""

import csv
import io
import logging

logger = logging.getLogger(__name__)

MIME_PDF = "application/pdf"
_MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_MIME_CSV = ("text/csv", "application/csv", "text/comma-separated-values")

# Cortes defensivos para no mandar planillas enormes a la IA.
_MAX_FILAS = 200
_MAX_COLS = 30
_MAX_TEXTO = 20_000  # caracteres


def _ext(filename: str | None) -> str:
    return (filename or "").rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""


def es_pdf(mime: str | None, filename: str | None) -> bool:
    return (mime or "").lower() == MIME_PDF or _ext(filename) == "pdf"


def es_planilla(mime: str | None, filename: str | None) -> bool:
    m = (mime or "").lower()
    return m == _MIME_XLSX or m in _MIME_CSV or _ext(filename) in ("xlsx", "csv")


def es_documento_util(mime: str | None, filename: str | None) -> bool:
    """True si es un adjunto que sabemos leer (PDF o planilla)."""
    return es_pdf(mime, filename) or es_planilla(mime, filename)


def _xlsx_a_texto(data: bytes) -> str | None:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - dependencia declarada en pyproject
        logger.warning("openpyxl no instalado; no se puede leer .xlsx")
        return None
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - archivo corrupto/no-xlsx: no cortar la ingesta
        logger.exception("No se pudo abrir el .xlsx")
        return None
    bloques: list[str] = []
    for ws in wb.worksheets:
        filas: list[str] = []
        for i, fila in enumerate(ws.iter_rows(values_only=True)):
            if i >= _MAX_FILAS:
                filas.append("… (planilla truncada)")
                break
            celdas = [
                "" if c is None else str(c).strip()
                for c in fila[:_MAX_COLS]
            ]
            if any(celdas):
                filas.append(" | ".join(celdas).rstrip(" |"))
        if filas:
            bloques.append(f"[Hoja: {ws.title}]\n" + "\n".join(filas))
    wb.close()
    texto = "\n\n".join(bloques).strip()
    return texto or None


def _csv_a_texto(data: bytes) -> str | None:
    for enc in ("utf-8-sig", "latin-1"):
        try:
            crudo = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        return None
    filas: list[str] = []
    for i, fila in enumerate(csv.reader(io.StringIO(crudo))):
        if i >= _MAX_FILAS:
            filas.append("… (planilla truncada)")
            break
        celdas = [c.strip() for c in fila[:_MAX_COLS]]
        if any(celdas):
            filas.append(" | ".join(celdas).rstrip(" |"))
    return "\n".join(filas).strip() or None


def planilla_a_texto(nombre: str | None, mime: str | None, data: bytes) -> str | None:
    """Convierte una planilla (Excel/CSV) a texto legible; None si no se pudo."""
    m = (mime or "").lower()
    if m == _MIME_XLSX or _ext(nombre) == "xlsx":
        texto = _xlsx_a_texto(data)
    elif m in _MIME_CSV or _ext(nombre) == "csv":
        texto = _csv_a_texto(data)
    else:
        return None
    if not texto:
        return None
    return texto[:_MAX_TEXTO]


def enriquecer_cuerpo(cuerpo: str, documentos: list[dict] | None) -> str:
    """Anexa al cuerpo el texto de las planillas adjuntas (Excel/CSV) para que la
    IA vea los ítems del pedido. Los PDF NO se anexan acá (van a la IA aparte)."""
    if not documentos:
        return cuerpo
    extras: list[str] = []
    for doc in documentos:
        if es_pdf(doc.get("mime"), doc.get("nombre")):
            continue
        texto = planilla_a_texto(doc.get("nombre"), doc.get("mime"), doc.get("data") or b"")
        if texto:
            nombre = doc.get("nombre") or "planilla"
            extras.append(f"\n\n--- Planilla adjunta: {nombre} ---\n{texto}")
    if not extras:
        return cuerpo
    return cuerpo + "".join(extras)
